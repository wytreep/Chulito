from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/home/ubuntu/Confirma_Entrega_Tecnica_2026-10-17/referencias')
DOCX = OUT / 'Aclaracion_Metodologica_PRD_v0.3.docx'
NAVY = '17365D'
BLUE = '246A8D'
RED = 'A02D23'
PALE = 'F3F7FA'
DISCLAIMER = 'S-0.1 es un escenario simulado, no resultados de entrevistas realizadas.'


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string('243447')
    normal.paragraph_format.space_after = Pt(7)

    header = sec.header.paragraphs[0]
    header.text = 'CONFIRMA  ·  ACLARACIÓN METODOLÓGICA DEL PRD v0.3'
    for r in header.runs:
        r.font.name = 'Arial'; r.font.size = Pt(8); r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(38)
    title.paragraph_format.space_after = Pt(5)
    r = title.add_run('Aclaración metodológica')
    r.font.name = 'Arial'; r.font.size = Pt(23); r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    r = sub.add_run('Complemento editorial para acompañar el PRD Confirma v0.3')
    r.font.name = 'Arial'; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(BLUE)

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for k,v in [('Documento relacionado','PRD CNF-PRD-001 v0.3'),('Complemento','CNF-NAM-PRD-001 · v0.1'),('Fecha','17 de octubre de 2026'),('Estatus','Aclaración editorial; no modifica requisitos ni decisiones del PRD')]:
        cells = table.add_row().cells
        cells[0].text = k; cells[1].text = v
        cells[0].width = Inches(1.75); cells[1].width = Inches(5.0)
        shade(cells[0], 'EAF1F6')
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for rr in p.runs:
                    rr.font.name='Arial'; rr.font.size=Pt(9)
        for rr in cells[0].paragraphs[0].runs:
            rr.font.bold=True; rr.font.color.rgb=RGBColor.from_string(NAVY)

    warning = doc.add_paragraph()
    warning.paragraph_format.space_before = Pt(18)
    warning.paragraph_format.space_after = Pt(14)
    warning.paragraph_format.left_indent = Inches(0.14)
    warning.paragraph_format.right_indent = Inches(0.14)
    warning.paragraph_format.keep_together = True
    r = warning.add_run(DISCLAIMER)
    r.bold=True; r.font.name='Arial'; r.font.size=Pt(13)
    r.font.color.rgb=RGBColor.from_string(RED)

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    rr = h.add_run('Alcance de la aclaración')
    rr.bold=True; rr.font.name='Arial'; rr.font.size=Pt(13); rr.font.color.rgb=RGBColor.from_string(NAVY)
    doc.add_paragraph('Las cifras y perfiles de S-0.1 son ficticios. No representan personas o negocios entrevistados y no constituyen evidencia de campo, mercado, preferencias, demanda ni validación del segmento o del alcance del MVP. En esta revisión se conservan únicamente como supuesto de trabajo para orientar preguntas y pruebas técnicas con datos sintéticos.')
    doc.add_paragraph('El PRD v0.3 original señala que S-0.1 es un “escenario de trabajo” y que las afirmaciones que dependan de él se revisan con entrevistas reales. Este complemento hace explícito que no se realizaron esas entrevistas. Las decisiones de producto deben contrastarse con evidencia real antes de aprobar o baselinar el MVP y antes de iniciar un piloto.')
    doc.add_paragraph('Esta página se agrega antes del PDF fuente solo para dar visibilidad al aviso en el paquete de revisión. El PRD v0.3 original que sigue se conserva sin alteraciones sustantivas; la presente nota no crea una nueva versión del PRD ni una declaración de conformidad normativa.')

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('Uso interno · Complemento editorial · 17-oct-2026')
    r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('687887')

    doc.core_properties.title = 'Aclaración metodológica del PRD Confirma v0.3'
    doc.core_properties.subject = DISCLAIMER
    doc.core_properties.comments = 'Anexo editorial separado; no altera el PRD fuente.'
    doc.save(DOCX)
    print(DOCX)

if __name__ == '__main__':
    main()
