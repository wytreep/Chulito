from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'Spike_Tecnico_Chulito_ST-01_v0.1.md'
OUT = ROOT / 'Spike_Tecnico_Chulito_ST-01_v0.1.docx'
NAVY = '1F3A5F'
BLUE = '315B86'
PALE = 'EAF0F6'
ALT = 'F4F7FA'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, text, end])


def add_inline(paragraph, text, base_size=None):
    tokens = re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1]); run.font.name = 'Consolas'; run.font.color.rgb = RGBColor.from_string(BLUE)
        else:
            # Render Markdown links as readable label (URL), without relying on HTML.
            link_pattern = re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)')
            cursor = 0
            for match in link_pattern.finditer(token):
                if match.start() > cursor:
                    paragraph.add_run(token[cursor:match.start()])
                paragraph.add_run(f'{match.group(1)} ({match.group(2)})')
                cursor = match.end()
            if cursor < len(token):
                paragraph.add_run(token[cursor:])
    if base_size:
        for run in paragraph.runs:
            run.font.size = Pt(base_size)


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if cells and all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in cells):
            continue
        rows.append(cells)
    return rows


def add_markdown_table(doc, lines):
    rows = parse_table(lines)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        tr_pr.append(cant_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            tr_pr.append(repeat)
    widths = {
        2: [1.75, 5.05],
        3: [1.55, 2.1, 3.15],
        4: [1.35, 1.55, 1.95, 1.95],
    }.get(cols)
    if widths:
        for i, width in enumerate(widths):
            for cell in table.columns[i].cells:
                cell.width = Inches(width)
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ci < len(row):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                add_inline(p, row[ci], base_size=8.1)
            if ri == 0:
                shade(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif ri % 2 == 0:
                shade(cell, ALT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        header = section.header.paragraphs[0]
        header.text = 'CONFIRMA  ·  SPIKE TÉCNICO ST-01'
        header.style = doc.styles['Normal']
        header.paragraph_format.space_after = Pt(3)
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Uso interno · Documento de trabajo    Página ')
        add_field(footer, 'PAGE')
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(105, 110, 118)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(38, 46, 56)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 10.5)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(9 if name == 'Heading 1' else 6)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True


def add_cover(doc, lines):
    doc.add_paragraph('CONFIRMA · PROYECTO DE SOFTWARE').runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    title = lines[0].lstrip('# ').replace('Chulito — ', '')
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(44)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(title)
    run.bold = True; run.font.name = 'Arial'; run.font.size = Pt(24); run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = lines[1].lstrip('# ').strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run(subtitle); run.font.size = Pt(16); run.font.color.rgb = RGBColor(88, 101, 115)
    metadata = []
    for line in lines[2:]:
        if line.startswith('**'):
            clean = line.replace('**', '').strip()
            if ':' in clean:
                k, v = clean.split(':', 1)
                metadata.append((k.strip(), v.strip()))
    if metadata:
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for key, value in metadata:
            cells = table.add_row().cells
            set_cell_margins(cells[0], 100, 100, 100, 100); set_cell_margins(cells[1], 100, 100, 100, 100)
            cells[0].text = key; cells[1].text = value
            shade(cells[0], PALE)
            for run in cells[0].paragraphs[0].runs:
                run.bold = True; run.font.color.rgb = RGBColor.from_string(NAVY)
            for cell in cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.add_run('Producto y requisitos · Documento de trabajo').italic = True
    doc.add_page_break()


def build():
    text = SOURCE.read_text(encoding='utf-8')
    lines = text.splitlines()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = 'Chulito — Spike técnico ST-01: canales de captura y modelo canónico'
    doc.core_properties.subject = 'Comparación de canales, esquema canónico, idempotencia y conciliación sintética'
    doc.core_properties.author = 'Proyecto Chulito'
    add_cover(doc, lines[:7])

    # Start at the first body section, after the cover metadata.
    start = next((i for i, line in enumerate(lines) if line.startswith('## 1. Resumen ejecutivo')), len(lines))
    i = start
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == '---':
            i += 1; continue
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F1F4F7'); pPr.append(shd)
            r = p.add_run('\n'.join(code)); r.font.name = 'Consolas'; r.font.size = Pt(8)
            i += 1; continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i]); i += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith('### '):
            p = doc.add_paragraph(style='Heading 2'); add_inline(p, line[4:].strip())
        elif line.startswith('## '):
            p = doc.add_paragraph(style='Heading 1'); add_inline(p, line[3:].strip())
        elif line.startswith('# '):
            p = doc.add_paragraph(style='Heading 1'); add_inline(p, line[2:].strip())
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), PALE); pPr.append(shd)
            add_inline(p, line[2:].strip())
        elif re.match(r'^\s*-\s+', line):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, re.sub(r'^\s*-\s+', '', line))
        elif re.match(r'^\s*\d+\.\s+', line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            add_inline(p, line.strip())
        else:
            p = doc.add_paragraph(); add_inline(p, line.strip())
        i += 1

    doc.save(OUT)
    print(f'created={OUT} bytes={OUT.stat().st_size}')


if __name__ == '__main__':
    build()
