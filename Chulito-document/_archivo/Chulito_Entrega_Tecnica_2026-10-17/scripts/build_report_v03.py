from pathlib import Path
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
import markdown
import re

ROOT = Path('/home/ubuntu/confirma-sintesis-documental')
MD_PATH = ROOT / 'Sintesis_Documental_Chulito_SD-0.3.md'
DOCX_PATH = ROOT / 'Sintesis_Documental_Chulito_SD-0.3.docx'
NAVY = '193654'
BLUE = '1F5A85'
TEAL = '2E7184'
PALE = 'EAF1F6'
PALE2 = 'F4F7FA'
TEXT = '243447'
MUTED = '5B6875'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=75, start=90, bottom=75, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for tag, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{tag}'))
        if node is None:
            node = OxmlElement(f'w:{tag}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:cantSplit')
    trPr.append(node)


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:tblHeader')
    node.set(qn('w:val'), 'true')
    trPr.append(node)


def set_paragraph_border(paragraph, side='bottom', color=BLUE, size='8', space='1'):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), size)
    el.set(qn('w:space'), space)
    el.set(qn('w:color'), color)
    pBdr.append(el)


def add_field(run, field_name):
    r = run._r
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' {field_name} '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for node in [begin, instr, separate, text, end]:
        r.append(node)


def style_run(run, *, bold=False, italic=False, color=TEXT, size=10.3, font='Arial', underline=False):
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is not None:
        rFonts.set(qn('w:eastAsia'), font)


def append_inline(paragraph, node, *, bold=False, italic=False, color=TEXT, size=10.3):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            style_run(run, bold=bold, italic=italic, color=color, size=size)
        return
    if not isinstance(node, Tag):
        return
    tag = node.name.lower()
    next_bold = bold or tag in ('strong', 'b')
    next_italic = italic or tag in ('em', 'i')
    next_color = BLUE if tag == 'a' else color
    for child in node.children:
        if tag == 'br':
            paragraph.add_run().add_break()
        else:
            append_inline(paragraph, child, bold=next_bold, italic=next_italic, color=next_color, size=size)


def add_text_to_paragraph(paragraph, element, size=10.3, color=TEXT):
    for child in element.children:
        append_inline(paragraph, child, size=size, color=color)


def add_callout(doc, element):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), PALE)
    pPr.append(shd)
    set_paragraph_border(p, side='left', color=TEAL, size='20', space='5')
    for child in element.children:
        append_inline(p, child, size=9.7, color=NAVY)
    return p


def add_table(doc, element, cover=False):
    rows = []
    thead = element.find('thead', recursive=False)
    tbody = element.find('tbody', recursive=False)
    if thead:
        for tr in thead.find_all('tr', recursive=False):
            rows.append(tr)
    if tbody:
        rows.extend(tbody.find_all('tr', recursive=False))
    if not rows:
        rows = element.find_all('tr')
    if not rows:
        return
    ncols = max(len(r.find_all(['th', 'td'], recursive=False)) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, tr in enumerate(rows):
        cells_html = tr.find_all(['th', 'td'], recursive=False)
        row = table.add_row()
        prevent_row_split(row)
        for ci in range(ncols):
            cell = row.cells[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pad = 140 if cover else 55
            set_cell_margins(cell, top=pad, start=70, bottom=pad, end=70)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            if ci < len(cells_html):
                add_text_to_paragraph(p, cells_html[ci], size=8.1 if ncols >= 4 else 8.6,
                                      color=WHITE if ri == 0 else TEXT)
            if ri == 0:
                set_cell_shading(cell, NAVY)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
            elif ri % 2 == 0:
                set_cell_shading(cell, PALE2)
        if ri == 0:
            repeat_table_header(row)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    return table


def add_list(doc, list_tag, depth=0):
    ordered = list_tag.name.lower() == 'ol'
    items = list_tag.find_all('li', recursive=False)
    for item_index, li in enumerate(items, start=1):
        if ordered:
            # Use explicit numbering so separate Markdown lists do not inherit Word's counter.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28 + 0.18 * depth)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            number = p.add_run(f'{item_index}. ')
            style_run(number, bold=True, color=BLUE, size=10.1)
        else:
            style = 'List Bullet 2' if depth > 0 else 'List Bullet'
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25 + 0.2 * depth)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        for child in li.children:
            if isinstance(child, Tag) and child.name.lower() in ('ul', 'ol'):
                continue
            if isinstance(child, Tag) and child.name.lower() == 'p':
                add_text_to_paragraph(p, child)
            else:
                append_inline(p, child)
        for nested in li.find_all(['ul', 'ol'], recursive=False):
            add_list(doc, nested, depth + 1)


def process_chunk(doc, html, chunk_index):
    soup = BeautifulSoup(html, 'html.parser')
    first_heading = True
    table_count = 0
    for element in soup.contents:
        if isinstance(element, NavigableString) and not str(element).strip():
            continue
        if not isinstance(element, Tag):
            continue
        tag = element.name.lower()
        if tag in ('h1', 'h2', 'h3', 'h4'):
            text = element.get_text(' ', strip=True)
            if chunk_index == 0 and first_heading and tag == 'h1':
                p = doc.add_paragraph()
                p.style = doc.styles['Title']
                p.paragraph_format.space_before = Pt(50)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                style_run(run, bold=True, color=NAVY, size=28, font='Arial')
            else:
                level = 1 if tag in ('h1', 'h2') else 2
                p = doc.add_heading(text, level=level)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.keep_together = True
                for run in p.runs:
                    style_run(run, bold=True, color=BLUE if level == 1 else TEAL,
                              size=17 if level == 1 else 12.5)
            first_heading = False
        elif tag == 'p':
            text = element.get_text(' ', strip=True)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.12
            if chunk_index == 0 and text.startswith('Consolidación documental'):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for child in element.children:
                    append_inline(p, child, size=14, color=MUTED)
                p.paragraph_format.space_after = Pt(20)
            elif text.startswith('**'):
                add_text_to_paragraph(p, element)
            else:
                add_text_to_paragraph(p, element)
        elif tag == 'blockquote':
            add_callout(doc, element)
        elif tag == 'table':
            add_table(doc, element, cover=(chunk_index == 0 and table_count == 0))
            table_count += 1
        elif tag in ('ul', 'ol'):
            add_list(doc, element)
        elif tag == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            set_paragraph_border(p, side='bottom', color='B8C7D3', size='5', space='1')
        elif tag == 'pre':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(5)
            r = p.add_run(element.get_text())
            style_run(r, font='Consolas', size=9, color=TEXT)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [
        ('Heading 1', 17, BLUE, 14, 6),
        ('Heading 2', 12.5, TEAL, 10, 4),
        ('Heading 3', 11, NAVY, 8, 3),
    ]:
        s = doc.styles[name]
        s.font.name = 'Arial'
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    # Running header
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(3)
    r = hp.add_run('CONFIRMA  ·  SÍNTESIS DOCUMENTAL')
    style_run(r, bold=True, color=NAVY, size=8.2)
    r = hp.add_run(' ' * 4 + 'CNF-SD-001  ·  SD-0.3')
    style_run(r, color=MUTED, size=8.2)
    set_paragraph_border(hp, side='bottom', color='B8C7D3', size='5', space='2')
    notice = header.add_paragraph()
    notice.paragraph_format.space_after = Pt(1)
    nr = notice.add_run('S-0.1 es un escenario simulado, no resultados de entrevistas realizadas.')
    style_run(nr, bold=True, color='A02D23', size=7.2)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(2)
    r = fp.add_run('Interno — borrador de trabajo  ·  Página ')
    style_run(r, color=MUTED, size=8)
    r = fp.add_run()
    style_run(r, color=MUTED, size=8)
    add_field(r, 'PAGE')
    r = fp.add_run(' de ')
    style_run(r, color=MUTED, size=8)
    r = fp.add_run()
    style_run(r, color=MUTED, size=8)
    add_field(r, 'NUMPAGES')

    props = doc.core_properties
    props.title = 'Chulito — Síntesis documental SD-0.3: supuestos de trabajo y pendientes'
    props.subject = 'S-0.1 identificado como escenario simulado; síntesis documental y supuestos de trabajo'
    props.author = 'Edwin Carabalí'
    props.keywords = 'Chulito, síntesis documental, PRD, SRS, requisitos, Fase 0'


def main():
    raw = MD_PATH.read_text(encoding='utf-8')
    chunks = re.split(r'\n\s*<!-- PAGEBREAK -->\s*\n', raw)
    doc = Document()
    configure_document(doc)
    for i, chunk in enumerate(chunks):
        if i:
            doc.add_page_break()
        html = markdown.markdown(chunk, extensions=['tables', 'sane_lists'])
        process_chunk(doc, html, i)
    doc.save(DOCX_PATH)
    print(f'created={DOCX_PATH} bytes={DOCX_PATH.stat().st_size} chunks={len(chunks)}')


if __name__ == '__main__':
    main()
