from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.table import _Row
from docx.shared import Pt, RGBColor

SRC = Path('/home/ubuntu/confirma-srs-v02-profesional/SRS_Chulito_v0.2_29148_alineacion_estructural.docx')
OUT_DIR = Path('/home/ubuntu/confirma-srs-v03-profesional')
OUT = OUT_DIR / 'SRS_Chulito_v0.3_29148_alineacion_estructural.docx'
DATE = '17 de octubre de 2026'
DISCLAIMER = 'S-0.1 es un escenario simulado, no resultados de entrevistas realizadas.'


def set_paragraph_text(paragraph, text):
    runs = paragraph.runs
    rpr = deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    paragraph.clear()
    new_run = paragraph.add_run(text)
    if rpr is not None:
        new_run._r.insert(0, rpr)


def set_cell_text(cell, text):
    p = cell.paragraphs[0]
    set_paragraph_text(p, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def replace_all(part, old, new):
    for p in part.paragraphs:
        if old in p.text:
            set_paragraph_text(p, p.text.replace(old, new))
    for table in part.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        set_paragraph_text(p, p.text.replace(old, new))


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)
    doc.core_properties.title = 'Especificación de requisitos de software (SRS) — Chulito v0.3'
    doc.core_properties.subject = 'SRS v0.3 trazado al PRD v0.3; S-0.1 identificado como escenario simulado; revisión de pendientes de implementación y validación'
    doc.core_properties.keywords = 'Chulito, SRS, PRD v0.3, escenario simulado, requisitos, trazabilidad, idempotencia, NestJS, TypeScript'
    doc.core_properties.comments = 'Revisión metodológica: S-0.1 no son entrevistas; datos ficticios como supuesto de trabajo. ADR-0001 de compatibilidad y port a TypeScript/NestJS pendientes.'

    # Version, date and status on cover and document-control metadata.
    for ti in (0, 1):
        for row in doc.tables[ti].rows:
            label = ' '.join(row.cells[0].text.split()).rstrip(':')
            if label in ('Versión', 'Versión vigente'):
                set_cell_text(row.cells[1], '0.3')
            elif label == 'Fecha':
                set_cell_text(row.cells[1], DATE)
            elif label == 'Edición de formato':
                set_cell_text(row.cells[1], DATE)
            elif label == 'Estado':
                set_cell_text(row.cells[1], 'Borrador v0.3; derivado del PRD v0.3; pendiente de revisión y aprobación; S-0.1 es escenario simulado.')
            elif label == 'Documento de origen':
                set_cell_text(row.cells[1], 'PRD CNF-PRD-001 v0.3 (12-oct-2026); S-0.1 identificado como escenario simulado, no evidencia de campo.')

    previous_v02 = doc.tables[2].rows[3]
    set_cell_text(previous_v02.cells[3], 'Actualizó origen al PRD v0.3 y corrigió trazas RF-06; la revisión v0.3 aclara que S-0.1 es simulado, no evidencia de entrevistas.')
    history = doc.tables[2].add_row()
    for cell, value in zip(history.cells, [
        '0.3', DATE, 'Edwin Carabalí',
        'Aclaración explícita de S-0.1 como escenario simulado; se registran la validación de producto, el spike de compatibilidad ADR-0001 y el port a TypeScript/NestJS como pendientes.'
    ]):
        set_cell_text(cell, value)

    # Cover summary and data provenance.
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('ESTADO DEL CONTENIDO ·'):
            set_paragraph_text(p, 'ESTADO DEL CONTENIDO · Borrador derivado del PRD v0.3. ' + DISCLAIMER + ' Sus cifras son ficticias y se usan solo como supuesto de trabajo para el diseño sintético; no validan segmento, problema ni alcance. Las marcas [HIP] y TBD identifican aspectos sujetos a validación o decisión.')
        elif t.startswith('Origen de los datos:'):
            set_paragraph_text(p, 'Origen de los datos: este SRS deriva del PRD v0.3. ' + DISCLAIMER + ' El escenario contiene cifras ficticias y no representa entrevistas ni preferencias reales. Se conserva como supuesto de trabajo para ejemplos y diseño técnico; la selección de segmento y alcance requiere evidencia de campo y aprobación antes de baselinar el MVP.')
        elif t.startswith('1.4 Referencias.'):
            set_paragraph_text(p, '1.4 Referencias. PRD v0.3; Guía de entrevistas v0.1; Informe de hallazgos S-0.1 (escenario simulado; no resultados de entrevistas realizadas); ISO/IEC/IEEE 29148:2018 (2.ª ed.); Ley 1581 de 2012.')
        elif t.startswith('1.5 Convenciones.'):
            set_paragraph_text(p, '1.5 Convenciones. Debe = obligatorio; debería = recomendable; podrá = capacidad opcional, no incluida necesariamente en el MVP. Prioridad: Alta, Media o Baja. Identificadores: SRS-F-xx (funcional), SRS-NFR-xx (no funcional), SRS-RN-xx (regla de negocio). [HIP] identifica hipótesis o metas provisionales. ' + DISCLAIMER + ' Los supuestos no equivalen a evidencia de usuario. Los pendientes se marcan TBD-n y se listan en la sección 10.')
        elif t.startswith('2.4 Restricciones.'):
            set_paragraph_text(p, t + ' El stack objetivo es NestJS con TypeScript; la compatibilidad de versiones de Node.js, TypeScript, NestJS y dependencias debe verificarse en el spike pendiente asociado a ADR-0001. La vertical slice actual está implementada en Python como especificación ejecutable, no como código de producción; su port a TypeScript está pendiente.')

    # Add explicit, visible simulation notice to the cover/status paragraph and to each running header.
    for section in doc.sections:
        for part in (section.header, section.first_page_header):
            replace_all(part, 'v0.2', 'v0.3')
            if not any(DISCLAIMER in p.text for p in part.paragraphs):
                p = part.add_paragraph()
                r = p.add_run(DISCLAIMER)
                r.bold = True
                r.font.name = 'Arial'
                r.font.size = Pt(7.2)
                r.font.color.rgb = RGBColor(160, 45, 35)
        for part in (section.footer, section.first_page_footer):
            replace_all(part, 'v0.2', 'v0.3')

    # Keep source evidence and technical implementation boundaries explicit.
    open_items = doc.tables[16]
    for row in open_items.rows[1:]:
        if row.cells[0].text.strip() == 'TBD-2':
            set_cell_text(row.cells[2], 'Validar uso/prioridad con investigación real y obtener muestras autorizadas/anonimizadas antes de integración real; S-0.1 no es evidencia de preferencias.')
    new_rows = [
        ['TBD-7', 'Validación real de segmento, problema y alcance del MVP (S-0.1 es simulado)', 'Entrevistas reales con perfiles definidos; síntesis de evidencia y aprobación del PRD antes de baselinar producto/piloto.'],
        ['TBD-8', 'Compatibilidad del stack NestJS/TypeScript/Node.js para ADR-0001', 'Spike separado de versiones, dependencias, toolchain y CI; ST-01 no cubre esta compatibilidad.'],
        ['TBD-9', 'Port de la vertical slice Python al stack objetivo', 'Traducir contrato, reglas y pruebas a TypeScript/NestJS; mantener fixtures sintéticos y no usar Python como backend de producción.'],
    ]
    for values in new_rows:
        row = open_items.add_row()
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    # Update any outdated reference to the previous content version in headers/footers.
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            replace_all(part, 'v0.2', 'v0.3')

    doc.save(OUT)
    print(f'created={OUT} bytes={OUT.stat().st_size}')
    print('version=0.3 disclaimer=visible+running-header pending=TBD-7,8,9')


if __name__ == '__main__':
    main()
